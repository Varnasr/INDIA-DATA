#!/usr/bin/env python3
"""Writes the review workbook: every string on the site with its Hindi, Urdu,
Telugu, Tamil and Bengali translations side by side, for a native reader to check.

    python3 scripts/i18n_workbook.py out.xlsx            # every language
    python3 scripts/i18n_workbook.py out.xlsx --lang hi  # one reviewer's copy

One sheet per language plus a combined sheet, or with --lang the one
language alone. Each row carries the pages the
string appears on and an empty column for the reviewer's correction, so the
corrections can be read back into i18n/<lang>.json key by key (then run
scripts/build_i18n.py). Needs openpyxl.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
NAMES = {'hi': 'Hindi', 'ur': 'Urdu', 'te': 'Telugu', 'ta': 'Tamil', 'bn': 'Bengali'}


def load(lang):
    return json.loads((ROOT / 'i18n' / (lang + '.json')).read_text(encoding='utf-8'))


def tier(pages):
    """1: the site's shell and the front page, which every reader meets.
    2: the first page of each place. 3: everything else."""
    if len(pages) >= 5 or '/' in pages:
        return 1
    if any(p.count('/') == 2 and not p.startswith('/assets') for p in pages):
        return 2
    return 3


def main(out, only=None):
    cat = json.loads((ROOT / 'i18n' / 'catalogue.json').read_text(encoding='utf-8'))
    langs = {only: NAMES[only]} if only else NAMES
    dicts = {k: load(k) for k in langs}
    keys = sorted(cat, key=lambda k: (tier(cat[k]['pages']), -len(cat[k]['pages']), min(cat[k]['pages']), k.lower()))
    wb = Workbook()
    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='8B1E3F')
    wrap = Alignment(wrap_text=True, vertical='top')

    def sheet(ws, cols, rows, rtl_cols=()):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        widths = {'No.': 6, 'Priority': 8, 'Pages': 22, 'Kind': 9, 'English': 60}
        for i, name in enumerate(cols, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = widths.get(name, 60 if name in NAMES.values() else 40)
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right' if name in rtl_cols else None,
                                           readingOrder=2 if name in rtl_cols else 0)
        ws.freeze_panes = 'F2'

    first = True
    if not only:
        ws = wb.active
        ws.title = 'All languages'
        sheet(ws, ['No.', 'Priority', 'Pages', 'Kind', 'English'] + list(NAMES.values()),
              [[i + 1, tier(cat[k]['pages']), ', '.join(cat[k]['pages']), cat[k]['kind'], k] + [dicts[l].get(k, '') for l in NAMES]
               for i, k in enumerate(keys)],
              rtl_cols=('Urdu',))
        first = False
    for lang, name in langs.items():
        ws = wb.active if first else wb.create_sheet(name)
        ws.title, first = name, False
        sheet(ws, ['No.', 'Priority', 'Pages', 'Kind', 'English', name, 'Correction', 'Reviewer note'],
              [[i + 1, tier(cat[k]['pages']), ', '.join(cat[k]['pages']), cat[k]['kind'], k, dicts[lang].get(k, ''), '', '']
               for i, k in enumerate(keys)],
              rtl_cols=(name, 'Correction') if lang == 'ur' else ())
    notes = wb.create_sheet('How to review', 0)
    for line in [
        'How to review these translations',
        '',
        'Each language has its own sheet. Read the English and the translation; if the translation is wrong, unclear or',
        'reads badly, write the corrected text in the Correction column and, if useful, a note. Leave Correction empty',
        'where the translation is right.',
        '',
        'Text in curly braces, such as {total} or {seat}, is filled in by the page with a figure or a name. Keep every',
        'one exactly as written, in Latin letters, but move it to wherever it belongs in the sentence.',
        'Tags such as <a href="...">...</a> or <b>...</b> are links and bold text. Keep them, with the text between',
        'them translated.',
        'Figures stay in Western digits with Indian grouping (6,62,063). Do not use an em dash.',
        'Names of places and parties that come from the data appear in charts and tables in Latin letters. That is',
        'expected.',
        '',
        'Rows are in order of how many readers see them. Priority 1 is the menus, buttons and front page, which',
        'every reader meets; priority 2 is the first page of each state; priority 3 is everything else. If time is',
        'short, priority 1 alone is worth doing.',
        '',
        '"Kind" says where a string comes from: markup is fixed text on a page, script is a sentence the page builds',
        'from the data. The Pages column says where it appears.',
    ]:
        notes.append([line])
    notes.column_dimensions['A'].width = 120
    notes['A1'].font = Font(bold=True, size=14)
    wb.save(out)
    print('wrote %s: %d strings x %d languages' % (out, len(keys), len(langs)))


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    lang = sys.argv[sys.argv.index('--lang') + 1] if '--lang' in sys.argv else None
    if lang:
        args = [a for a in args if a != lang]
    main(args[0] if args else 'translations-for-review.xlsx', lang)
