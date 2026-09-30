"""Small helpers shared by the two Word reports: styles, tables, charts.

The reports follow the structure of the Sabar Institute's SIR reports: Part I
analyses the whole area, Part II gives one chapter per district or
constituency, and a single section says what the lists show and what they do
not. Every figure is computed from the published JSON when the report is built.
"""
import base64
from PIL import Image
import html as H
import io
import os
import re
from datetime import date

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1C, 0x28, 0x33)
MUTED = RGBColor(0x4D, 0x5B, 0x69)
NF = lambda x: '{:,}'.format(int(round(x))) if x is not None else '–'


def indian(x):
    """6,62,063: Indian digit grouping, as on the pages."""
    if x is None:
        return '–'
    s = str(int(round(x)))
    neg = s.startswith('-')
    s = s.lstrip('-')
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ','.join(parts) + ',' + tail
    return ('-' if neg else '') + s


def pct(x, k=1):
    return '–' if x is None else ('%.' + str(k) + 'f%%') % x


def lakh(x):
    return '%.1f lakh' % (x / 1e5)


class Report:
    def __init__(self, accent):
        self.doc = Document()
        self.accent = accent
        # The same report as print-ready HTML, written alongside the Word file
        # and printed to PDF by headless Chromium (scripts/reports/pdf.mjs).
        self.html = []
        self.heads = []
        self.figs = []
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ('left_margin', 'right_margin'):
            setattr(sec, side, Cm(2.2))
        sec.top_margin = sec.bottom_margin = Cm(2)
        st = self.doc.styles['Normal']
        st.font.name = 'Calibri'
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        st.font.size = Pt(10.5)
        st.font.color.rgb = INK
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.15
        for lvl, size in [(1, 17), (2, 13.5), (3, 11.5)]:
            h = self.doc.styles['Heading %d' % lvl]
            h.font.name = 'Calibri'
            h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = accent if lvl < 3 else INK
            h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(4)
        self.doc.core_properties.author = 'Varna Sri Raman'
        self.doc.core_properties.last_modified_by = 'Varna Sri Raman'
        self._footer()

    def _footer(self):
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('Varna Sri Raman · Totals only, never names or EPIC numbers · Page ')
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED
        for code in ('begin', None, 'end'):
            run = p.add_run()
            run.font.size = Pt(8)
            if code:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), code)
                run._r.append(el)
            else:
                el = OxmlElement('w:instrText')
                el.set(qn('xml:space'), 'preserve')
                el.text = 'PAGE'
                run._r.append(el)

    # ------------------------------------------------------------ text
    def cover(self, kicker, title, subtitle, lines):
        self.title = title
        self.html.append('<section class="cover"><p class="kicker">%s</p><h1 class="title">%s</h1><p class="subtitle">%s</p><dl>%s</dl><p class="by">Varna Sri Raman · %s</p></section>' % (
            H.escape(kicker), H.escape(title), H.escape(subtitle), ''.join('<div><dt>%s</dt><dd>%s</dd></div>' % (H.escape(v), H.escape(l)) for l, v in lines), date.today().strftime('%-d %B %Y')))
        d = self.doc
        for _ in range(5):
            d.add_paragraph()
        p = d.add_paragraph()
        r = p.add_run(kicker.upper())
        r.font.size, r.bold, r.font.color.rgb = Pt(10), True, self.accent
        p = d.add_paragraph()
        r = p.add_run(title)
        r.font.size, r.bold = Pt(28), True
        p = d.add_paragraph()
        r = p.add_run(subtitle)
        r.font.size, r.font.color.rgb = Pt(14), MUTED
        d.add_paragraph()
        for label, value in lines:
            p = d.add_paragraph()
            r = p.add_run(value + '  ')
            r.font.size, r.bold, r.font.color.rgb = Pt(18), True, self.accent
            r = p.add_run(label)
            r.font.size = Pt(11)
        d.add_paragraph()
        p = d.add_paragraph('Varna Sri Raman · ' + date.today().strftime('%-d %B %Y'))
        p.runs[0].font.color.rgb = MUTED
        self.page_break()

    def h1(self, t):
        self.doc.add_heading(t, 1)
        if t != 'Contents':
            i = len(self.heads); self.heads.append((1, t)); self.html.append('<h1 id="s%d">%s</h1>' % (i, H.escape(t)))

    def h2(self, t):
        self.doc.add_heading(t, 2)
        i = len(self.heads); self.heads.append((2, t)); self.html.append('<h2 id="s%d">%s</h2>' % (i, H.escape(t)))

    def h3(self, t):
        self.doc.add_heading(t, 3)
        self.html.append('<h3>%s</h3>' % H.escape(t))

    @staticmethod
    def _h(t):
        return re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', H.escape(t))

    def p(self, t, italic=False, size=None, color=None):
        self.html.append('<p%s>%s</p>' % (' class="note"' if italic else '', self._h(t)))
        para = self.doc.add_paragraph()
        self._rich(para, t, italic)
        if size or color:
            for r in para.runs:
                if size:
                    r.font.size = Pt(size)
                if color:
                    r.font.color.rgb = color
        return para

    def _rich(self, para, t, italic=False):
        # **bold** spans, nothing else.
        parts = t.split('**')
        for i, part in enumerate(parts):
            if part:
                r = para.add_run(part)
                r.bold = i % 2 == 1
                r.italic = italic

    def bullets(self, items):
        self.html.append('<ul>' + ''.join('<li>%s</li>' % self._h(x) for x in items) + '</ul>')
        for t in items:
            para = self.doc.add_paragraph(style='List Bullet')
            self._rich(para, t)

    def numbered(self, items):
        self.html.append('<ol>' + ''.join('<li>%s</li>' % self._h(x) for x in items) + '</ol>')
        for t in items:
            para = self.doc.add_paragraph(style='List Number')
            self._rich(para, t)

    def callout(self, title, text):
        self.html.append('<div class="callout"><b>%s</b><p>%s</p></div>' % (H.escape(title), self._h(text)))
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        self._shade(c, 'FFF6D9')
        self._border(c, left=self.accent)
        para = c.paragraphs[0]
        r = para.add_run(title)
        r.bold = True
        r.font.color.rgb = self.accent
        para2 = c.add_paragraph()
        self._rich(para2, text)
        self.doc.add_paragraph()

    def stats(self, items):
        """Headline figures in a row of boxes, as on the dashboards."""
        self.html.append('<div class="stats">' + ''.join('<div><b>%s</b><span>%s</span></div>' % (H.escape(v), H.escape(l)) for v, l in items) + '</div>')
        n = len(items)
        t = self.doc.add_table(rows=1, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (value, label) in enumerate(items):
            c = t.rows[0].cells[i]
            self._shade(c, 'F4F6F7')
            self._border(c, left=self.accent)
            para = c.paragraphs[0]
            r = para.add_run(value)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(15), self.accent
            p2 = c.add_paragraph()
            r2 = p2.add_run(label)
            r2.font.size, r2.font.color.rgb = Pt(8.5), MUTED
        self.doc.add_paragraph()

    def table(self, head, rows, widths=None, right=None, size=8.5):
        right = right if right is not None else set(range(1, len(head)))
        cls = lambda i: ' class="r"' if i in right else ''
        self.html.append('<table style="font-size:%.1fpt"><thead><tr>%s</tr></thead><tbody>%s</tbody></table>' % (size, ''.join('<th%s>%s</th>' % (cls(i), H.escape(h)) for i, h in enumerate(head)),
            ''.join('<tr>%s</tr>' % ''.join('<td%s>%s</td>' % (cls(i), H.escape(str(v))) for i, v in enumerate(r)) for r in rows)))
        t = self.doc.add_table(rows=1, cols=len(head))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]
            self._shade(c, self._hex(self.accent))
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(size), RGBColor(0xFF, 0xFF, 0xFF)
            if i in right:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for k, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                para = cells[i].paragraphs[0]
                r = para.add_run(str(v))
                r.font.size = Pt(size)
                if i in right:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if k % 2:
                    self._shade(cells[i], 'F4F6F7')
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self._repeat_header(t)
        self.doc.add_paragraph()
        return t

    def figure(self, fig, caption, width=16):
        buf = io.BytesIO()
        fig.savefig(buf, format='png', dpi=200, bbox_inches='tight')
        plt.close(fig)
        buf.seek(0)
        # A marker for now: the print copy embeds the image, the web copy
        # links a file beside the page (save_html, save_web).
        self.figs.append((buf.getvalue(), width, caption))
        self.html.append('<!--FIG%d-->' % (len(self.figs) - 1))
        buf.seek(0)
        self.doc.add_picture(buf, width=Cm(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para = self.doc.add_paragraph()
        r = para.add_run(caption)
        r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), MUTED

    def page_break(self):
        self.html.append('<div class="pb"></div>')
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self):
        self.html.append('<!--TOC-->')
        para = self.doc.add_paragraph()
        for code, text in (('begin', None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ('separate', None), (None, None), ('end', None)):
            run = para.add_run()
            if code:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), code)
                run._r.append(el)
            elif text:
                el = OxmlElement('w:instrText')
                el.set(qn('xml:space'), 'preserve')
                el.text = text
                run._r.append(el)
            else:
                run.text = 'Right-click and choose Update Field to show the contents.'
                run.italic = True
        self.page_break()

    def save(self, path):
        # Ask Word to refresh the contents field when the file is opened.
        settings = self.doc.settings.element
        el = OxmlElement('w:updateFields')
        el.set(qn('w:val'), 'true')
        settings.append(el)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.doc.save(path)

    def save_html(self, path):
        toc = '<section class="toc"><h1>Contents</h1><ol>' + ''.join('<li class="l%d"><a href="#s%d">%s</a></li>' % (lvl, i, H.escape(t)) for i, (lvl, t) in enumerate(self.heads)) + '</ol></section><div class="pb"></div>'
        body = ''.join(self.html).replace('<!--TOC-->', toc)
        # One break between pages: the cover already breaks after itself, and
        # the contents carry their own, so repeated markers would print blank pages.
        body = re.sub(r'(<div class="pb"></div>)+', '<div class="pb"></div>', body)
        body = re.sub(r'(<section class="cover">.*?</section>)<div class="pb"></div>', r'\1', body, flags=re.S)
        for i, (png, width, cap) in enumerate(self.figs):
            body = body.replace('<!--FIG%d-->' % i, '<figure><img src="data:image/png;base64,%s" style="width:%.1fcm" alt="%s"><figcaption>%s</figcaption></figure>' % (base64.b64encode(png).decode(), width, H.escape(cap), H.escape(cap)))
        a = self._hex(self.accent)
        css = CSS.replace('ACCENT', '#' + a)
        doc = '<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>%s</title><style>%s</style></head><body>%s</body></html>' % (H.escape(self.title), css, body)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        open(path, 'w', encoding='utf-8').write(doc)

    def save_web(self, folder):
        """The report as a fragment for a page on the site: the contents list,
        the body, and each figure as a PNG beside the page. Tables go in a
        scrolling wrapper so a phone can read them. scripts/reports/web.py
        puts the fragment into the site's page shell."""
        os.makedirs(folder, exist_ok=True)
        for f in os.listdir(folder):
            if f.startswith('fig-') and f.endswith('.png'):
                os.remove(os.path.join(folder, f))
        toc = '<nav class="toc" aria-labelledby="toc-h"><h2 id="toc-h">Contents</h2><ol>' + ''.join('<li class="l%d"><a href="#s%d">%s</a></li>' % (lvl, i, H.escape(t)) for i, (lvl, t) in enumerate(self.heads)) + '</ol></nav>'
        body = ''.join(self.html).replace('<!--TOC-->', toc).replace('<div class="pb"></div>', '')
        for i, (png, width, cap) in enumerate(self.figs):
            name = 'fig-%02d.png' % (i + 1)
            # The print copy is 200 dpi; a screen needs no more than 1,400px.
            im = Image.open(io.BytesIO(png))
            if im.width > 1400:
                im = im.resize((1400, round(im.height * 1400 / im.width)), Image.LANCZOS)
            im.save(os.path.join(folder, name), optimize=True)
            body = body.replace('<!--FIG%d-->' % i, '<figure><img src="%s" alt="%s" loading="lazy"><figcaption>%s</figcaption></figure>' % (name, H.escape(cap), H.escape(cap)))
        body = body.replace('<table', '<div class="twrap" tabindex="0"><table').replace('</table>', '</table></div>')
        open(os.path.join(folder, 'report.fragment.html'), 'w', encoding='utf-8').write(body)

    # ------------------------------------------------------------ plumbing
    @staticmethod
    def _hex(rgb):
        return '%02X%02X%02X' % (rgb[0], rgb[1], rgb[2])

    @staticmethod
    def _shade(cell, hexfill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hexfill)
        tcPr.append(shd)

    def _border(self, cell, left):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement('w:tcBorders')
        el = OxmlElement('w:left')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '24')
        el.set(qn('w:color'), self._hex(left))
        b.append(el)
        tcPr.append(b)

    @staticmethod
    def _repeat_header(t):
        trPr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement('w:tblHeader')
        el.set(qn('w:val'), 'true')
        trPr.append(el)


def style_axes(ax, title=None):
    ax.spines[['top', 'right']].set_visible(False)
    ax.tick_params(labelsize=8)
    ax.grid(axis='x' if ax.get_xscale() == 'log' else 'y', color='#E1E5E8', linewidth=0.6)
    ax.set_axisbelow(True)
    if title:
        ax.set_title(title, fontsize=10, loc='left', fontweight='bold')


CSS = """
@page { size: A4; margin: 18mm 18mm 20mm; }
body { font-family: Calibri, Carlito, 'Segoe UI', 'Noto Sans', sans-serif; color: #1C2833; font-size: 10.5pt; line-height: 1.45; }
h1, h2, h3 { color: ACCENT; line-height: 1.25; break-after: avoid; }
h1 { font-size: 18pt; margin: 0 0 8pt; }
h2 { font-size: 13.5pt; margin: 16pt 0 4pt; }
h3 { font-size: 11.5pt; color: #1C2833; }
p { margin: 0 0 6pt; }
p.note { font-style: italic; font-size: 8.5pt; color: #4D5B69; }
.pb { break-after: page; }
.cover { padding-top: 60mm; break-after: page; }
.cover .kicker { font-weight: 700; color: ACCENT; letter-spacing: .08em; text-transform: uppercase; font-size: 10pt; }
.cover .title { font-size: 30pt; color: #1C2833; margin: 4pt 0; }
.cover .subtitle { font-size: 14pt; color: #4D5B69; }
.cover dl { margin: 18pt 0; }
.cover dl div { margin: 6pt 0; }
.cover dt { display: inline; font-size: 18pt; font-weight: 700; color: ACCENT; margin-right: 6pt; }
.cover dd { display: inline; margin: 0; font-size: 11pt; }
.cover .by { color: #4D5B69; margin-top: 24pt; }
.toc ol { list-style: none; padding: 0; }
.toc li.l1 { font-weight: 700; margin-top: 6pt; }
.toc li.l2 { margin-left: 12pt; }
.toc a { color: #1C2833; text-decoration: none; }
.stats { display: flex; gap: 6pt; margin: 6pt 0 10pt; break-inside: avoid; }
.stats div { flex: 1; background: #F4F6F7; border-left: 3pt solid ACCENT; padding: 6pt 8pt; }
.stats b { display: block; font-size: 15pt; color: ACCENT; }
.stats span { font-size: 8.5pt; color: #4D5B69; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 10pt; }
thead { display: table-header-group; }
th { background: ACCENT; color: #fff; text-align: left; padding: 3pt 4pt; }
td { padding: 2.5pt 4pt; border-bottom: 0.5pt solid #D5D8DC; }
tr { break-inside: avoid; }
tbody tr:nth-child(even) td { background: #F4F6F7; }
.r { text-align: right; }
figure { margin: 6pt 0 10pt; text-align: center; break-inside: avoid; }
figure img { max-width: 100%; }
figcaption { font-style: italic; font-size: 8.5pt; color: #4D5B69; text-align: left; }
ul, ol { margin: 0 0 8pt 16pt; padding: 0; }
li { margin-bottom: 3pt; }
.callout { background: #FFF6D9; border-left: 3pt solid ACCENT; padding: 6pt 8pt; margin: 6pt 0; }
"""
